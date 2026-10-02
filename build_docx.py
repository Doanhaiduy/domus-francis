import docx
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(148)
    section.page_height = Mm(210)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(16)
    section.right_margin = Mm(14)
    
    # Hide footer on first page (Cover page)
    section.different_first_page_header_footer = True
    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    for fp in first_footer.paragraphs:
        fp.text = ""
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)
        
    # Configure default Normal style
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.space_before = Pt(0)
    
    # Configure footer for pages 2 to 44
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_p.paragraph_format.space_before = Pt(0)
    f_p.paragraph_format.space_after = Pt(0)
    f_p.paragraph_format.line_spacing = 1.0
    
    # Add page number field to footer (ECMA-376 compliant: fldSimple as child of w:p)
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    r_field = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '19')
    rPr.append(sz)
    r_field.append(rPr)
    fldSimple.append(r_field)
    f_p._p.append(fldSimple)
    
    f_run2 = f_p.add_run('\nLưu xá Phanxicô')
    f_run2.font.name = 'Times New Roman'
    f_run2.font.size = Pt(9)
    f_run2.font.italic = False
    
    return doc

def add_p(doc, text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=2, size=10, underline=False, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.underline = underline
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
    return p

def add_p_mixed(doc, runs, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=2, line_spacing=1.15):
    """
    runs is a list of tuples: (text, bold, italic, underline, size)
    default size 10
    """
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    for item in runs:
        text = item[0]
        bold = item[1] if len(item) > 1 else False
        italic = item[2] if len(item) > 2 else False
        underline = item[3] if len(item) > 3 else False
        size = item[4] if len(item) > 4 else 10
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.underline = underline
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
    return p

def add_heading_1(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=3, underline=False):
    p = doc.add_paragraph(text, style='Heading 1')
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12.5)
        r.bold = True
        r.underline = underline
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_2(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=5, space_after=2, underline=False):
    p = doc.add_paragraph(text, style='Heading 2')
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.bold = True
        r.underline = underline
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def build_all():
    doc = create_document()
    
    # =========================================================================
    # TRANG 1: BÌA TRƯỚC
    # =========================================================================
    add_p(doc, space_before=30)
    add_p(doc, "Lưu xá Phanxicô", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=70)
    add_p(doc, "Giờ Kinh Phụng Vụ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=20, space_after=10)
    add_p(doc, "(Kinh Tối)", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=15, space_after=120)
    
    doc.add_page_break() # -> TRANG 2
    
    # =========================================================================
    # TRANG 2: CẤU TRÚC & MỤC LỤC
    # =========================================================================
    add_p(doc, "Cấu trúc", bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=11, space_before=8, space_after=3)
    add_p(doc, "GIỜ KINH TỐI LƯU XÁ PHANXICÔ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11.5, space_after=2)
    add_p(doc, "(Nha Trang)", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=8)
    
    # Divider line
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(10)
    r_line = p_line.add_run("―" * 38)
    r_line.font.name = 'Times New Roman'
    r_line.font.size = Pt(9)
    
    def add_toc_line(doc, title, page_str, bold=False, indent=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2.5)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.15
        
        tab_stops = p.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Pt(325), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        
        prefix = "        " if indent else ""
        r1 = p.add_run(prefix + title)
        r1.bold = bold
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(9.5)
        
        r2 = p.add_run("\t" + page_str)
        r2.bold = bold
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(9.5)
        return p
    
    add_toc_line(doc, "GIỜ KINH PHỤNG VỤ - KINH TỐI", "3", bold=True)
    add_toc_line(doc, "Giáo đầu & Nghi thức Sám hối", "3", indent=True)
    add_toc_line(doc, "Thánh thi", "5", indent=True)
    add_toc_line(doc, "Thánh Vịnh các ngày trong tuần", "8", indent=True)
    add_toc_line(doc, "Thứ Bảy (8) • Chúa Nhật (10) • Thứ Hai (12)", "", indent=True)
    add_toc_line(doc, "Thứ Ba (14) • Thứ Tư (16) • Thứ Năm (18) • Thứ Sáu (19)", "", indent=True)
    add_toc_line(doc, "Lời Chúa", "21", indent=True)
    add_toc_line(doc, "Xướng đáp & Ca chúc tụng Tin Mừng", "23", indent=True)
    add_toc_line(doc, "Lời nguyện kết thúc các ngày", "24", indent=True)
    add_toc_line(doc, "LẦN CHUỖI KÍNH ĐỨC MẸ", "26", bold=True)
    add_toc_line(doc, "BÀI HÁT KÍNH ĐỨC MẸ & THÁNH PHANXICÔ", "28", bold=True)
    add_toc_line(doc, "KINH CẦU ÂN NHÂN & THÂN NHÂN", "40", bold=True)
    add_toc_line(doc, "PHỤ LỤC I: KINH NGUYỆN PHANXICÔ", "41", bold=True)
    add_toc_line(doc, "PHỤ LỤC II: CÁC CA VÃN KÍNH ĐỨC MẸ", "43", bold=True)
    
    doc.add_page_break() # -> TRANG 3
    
    # =========================================================================
    # TRANG 3: GIÁO ĐẦU & SÁM HỐI 1
    # =========================================================================
    add_heading_2(doc, "Giáo đầu", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=8)
    
    add_p(doc, "Chủ sự:", bold=True, italic=True, space_after=2, size=10.5)
    add_p(doc, "Lạy Chúa Trời, xin tới giúp con", space_after=6, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Cộng đoàn:", bold=True, space_after=2, size=10.5)
    add_p(doc, "Muôn lạy Chúa, xin mau phù trợ.", space_after=6, size=10.5, line_spacing=1.25)
    
    add_p(doc, "(Chủ sự) Vinh danh Chúa Cha và Chúa Con\nCùng vinh danh Thánh Thần Thiên Chúa\nTự muôn đời và chính hiện nay,\nluôn mãi đến thiên thu vạn đại. Amen. Ha-lê-lui-a.", space_after=3, size=10.5, line_spacing=1.25)
    add_p(doc, "(Mùa Chay bỏ Amen Ha-lê-lui-a)", italic=True, space_after=10, size=9.5)
    
    add_p(doc, "(Sau đó thinh lặng giây lát để xét mình. Có thể dùng công thức thống hối như trong thánh lễ)", italic=True, space_after=10, size=9.5)
    
    add_heading_2(doc, "SÁM HỐI", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3)
    add_p(doc, "(Chọn một trong các mẫu sám hối sau đây)", italic=True, space_after=6, size=9.5)
    
    add_p(doc, "Nghi Thức Sám Hối 1: Kinh Ăn Năn Tội", bold=True, italic=True, space_after=4, size=10.5)
    add_p(doc, "Lạy Chúa (con), Chúa là Đấng trọn tốt trọn lành vô cùng. Chúa đã dựng nên con, và cho Con Chúa ra đời, chịu nạn chịu chết vì con, mà con đã cả lòng phản nghịch lỗi nghĩa cùng Chúa, thì con lo buồn đau đớn, cùng chê ghét mọi tội con trên hết mọi sự; con dốc lòng chừa cải, và nhờ ơn Chúa, thì con sẽ lánh xa dịp tội, cùng làm việc đền tội cho xứng. Amen.", space_after=6, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 4
    
    # =========================================================================
    # TRANG 4: NGHI THỨC SÁM HỐI 2 & 3
    # =========================================================================
    add_p(doc, "Nghi Thức Sám Hối 2: Kinh Thú Nhận", bold=True, italic=True, space_before=6, space_after=3)
    add_p_mixed(doc, [
        ("Chủ sự: ", True, True),
        ("Chúng ta hãy nhìn nhận tội lỗi chúng ta, để xứng đáng cử hành giờ kinh tối này", False, False)
    ], space_after=4)
    
    add_p_mixed(doc, [
        ("Cộng đoàn: ", True, False),
        ("Tôi thú nhận cùng Thiên Chúa toàn năng, và cùng anh (chị) em: tôi đã phạm tội nhiều trong tư tưởng, lời nói, việc làm, và những điều thiếu sót. (đấm ngực và đọc). Lỗi tại tôi, lỗi tại tôi, lỗi tại tôi mọi đàng. (Rồi đọc tiếp). Vì vậy tôi xin Đức Bà Maria trọn đời đồng trinh, các Thiên Thần, các Thánh và anh (chị) em, khẩn cầu cho tôi trước tòa Thiên Chúa, Chúa chúng ta.", False, False)
    ], space_after=4)
    
    add_p_mixed(doc, [
        ("Chủ sự: ", True, True),
        ("Xin Thiên Chúa toàn năng thương xót, tha tội, và dẫn đưa chúng ta đến sự sống muôn đời.", False, True)
    ], space_after=2)
    
    add_p_mixed(doc, [
        ("Cộng đoàn: ", True, False),
        ("Amen.", False, True)
    ], space_after=8)
    
    add_p(doc, "Nghi Thức Sám Hối 3: Cầu Xin Lòng Thương Xót", bold=True, italic=True, space_after=3)
    add_p_mixed(doc, [
        ("Chủ sự: ", True, True),
        ("Chúng ta hãy nhìn nhận tội lỗi chúng ta, để xứng đáng cử hành giờ kinh tối này", False, False)
    ], space_after=2)
    add_p(doc, "(Thinh lặng giây lát)", italic=True, space_after=2)
    
    add_p_mixed(doc, [
        ("Chủ sự: ", True, True),
        ("Lạy Chúa, xin thương xót chúng con.", False, False)
    ], space_after=1.5)
    add_p_mixed(doc, [
        ("Cđ: ", True, False),
        ("Vì chúng con đã xúc phạm đến Chúa.", False, False)
    ], space_after=1.5)
    add_p_mixed(doc, [
        ("Chủ sự: ", True, True),
        ("Lạy Chúa, xin tỏ lòng từ bi Chúa cho chúng con.", False, False)
    ], space_after=1.5)
    add_p_mixed(doc, [
        ("Cđ: ", True, False),
        ("Và ban ơn cứu độ cho chúng con.", False, False)
    ], space_after=1.5)
    add_p_mixed(doc, [
        ("Chủ sự: ", True, True),
        ("Xin Thiên Chúa toàn năng thương xót, tha tội, và dẫn đưa chúng ta đến sự sống muôn đời.", False, True)
    ], space_after=1.5)
    add_p_mixed(doc, [
        ("Cđ: ", True, False),
        ("Amen.", False, False)
    ], space_after=4)
    
    doc.add_page_break() # -> TRANG 5
    
    # =========================================================================
    # TRANG 5: THÁNH THI - NGOÀI MÙA PHỤC SINH (BÀI 1)
    # =========================================================================
    add_heading_2(doc, "Thánh thi", align=WD_ALIGN_PARAGRAPH.LEFT, underline=True, space_before=12, space_after=4)
    add_p(doc, "Ngoài Mùa Phục Sinh", bold=True, space_after=3, size=10.5)
    add_p(doc, "(Chọn một trong hai bài sau đây)", italic=True, space_after=12, size=9.5)
    
    add_p(doc, "Bài 1:", bold=True, space_after=6, size=11)
    add_p(doc, "Đêm tối xuống dần trên cõi thế,\nĐoàn con chạy đến Chúa càn khôn,\nNgàn muôn ơn thánh xin đổ xuống\nGiữ gìn chúng con cả xác hồn.", space_after=12, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Mơ thấy Chúa trời: lòng nguyện ước\nThầm mong cảm nghiệm lúc ngủ ngon,\nVầng đông lấp ló chân trời thẳm\nSẽ hát mừng Ngài khúc nhặt khoan.\nBan xuống chuỗi ngày đầy sức sống\nBồi thêm sinh khí kẻo tàn phai,\nChập chờn bóng tối gieo sợ hãi,\nXin hãy đốt lên lửa sáng ngời.", space_after=12, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Đồng thanh ca tụng Cha hằng hữu\nVà Thánh Tử Ngài, Đấng Phục Sinh,\nThần Linh thánh ái, ơn Phù trợ,\nMuôn thuở ngàn đời mãi hiển vinh.", space_after=6, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 6
    
    # =========================================================================
    # TRANG 6: THÁNH THI - NGOÀI MÙA PHỤC SINH (BÀI 2)
    # =========================================================================
    add_p(doc, "Bài 2:", bold=True, space_before=12, space_after=6, size=11)
    add_p(doc, "Muôn lạy Chúa Kitô Ánh Sáng\nBừng lên cho khuất dạng đêm đen\nHào quang muôn thuở diệu huyền\nSoi đường tín hữu đi trên cõi đời.", space_after=10, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Cúi xin Đấng tuyệt vời thánh thiện\nLắng nghe lời khẩn nguyện nài van,\nThương ban giấc ngủ yên hàn\nđược kề bên Chúa an toàn thong dong", space_after=10, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Dẫu mắt ngủ nhưng lòng vẫn thức,\nVẫn tin yêu một mực chân tình\nXin giơ tay hữu uy linh\nNhư đồn bảo vệ, như thành chở che.", space_after=10, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Cúi xin Đấng phù trì đoái đến\nNgăn chước thù độc hiểm gớm ghê,\nGiữ đoàn con cả đôi bề\nMáu Ngài tuôn đổ chuộc về thuở xưa.\nChúa Kitô là Vua nhân ái\nKính dâng Ngài cùng với Chúa Cha,\nHiệp cùng Thiên Chúa Ngôi Ba\nNgàn muôn phước cả vinh hoa đời đời.", space_after=6, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 7
    
    # =========================================================================
    # TRANG 7: THÁNH THI - MÙA PHỤC SINH
    # =========================================================================
    add_heading_2(doc, "Mùa Phục Sinh", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=8)
    add_p(doc, "Ngôi Lời Thánh Phụ quang vinh,\nNgài là chính Vị Cứu Tinh gian trần,\nNguyện cầu Ánh Sáng muôn dân\nCanh phòng giấc ngủ cho đoàn con đây.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Ngài ra định luật an bài\nNăm canh sáu khắc đổi thay tuyệt vời\nChúng con mệt mỏi rã rời\nMột đêm an nghỉ xin bồi dưỡng cho.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Âm ty Ngài đã phá xưa,\nNay xin cứu khỏi ác thù hiểm nguy,\nKẻo đàn con Chúa hư đi,\nUổng công giá máu chuộc về bấy nay.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Vậy dẫu cho xác thân này\nĐêm khuya giấc điệp mê say nặng nề.\nThì lòng trí vẫn một bề\nCùng Ngài gắn bó như quỳ hướng dương.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Ngợi khen Vua Cả thiên đường\nLừng danh đánh bại ma vương quỷ thần,\nNgàn đời hiển trị muôn dân\nCùng Ngôi Thánh Phụ Thánh Thần quang vinh.", space_after=6, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 8
    
    # =========================================================================
    # TRANG 8: THÁNH VỊNH THỨ BẢY - Tv 4 (Lời tạ ơn)
    # =========================================================================
    add_heading_1(doc, "Thánh Vịnh", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=4)
    add_p(doc, "THỨ BẢY", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11, space_after=2)
    add_p(doc, "(Tức sau kinh chiều I Chúa nhật, sau Kinh Chiều I các ngày lễ trọng)", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9, space_after=6)
    
    add_p(doc, "Tv 4 - Lời tạ ơn", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p(doc, "(Thiên Chúa biệt đãi Đức Kitô, Đấng Thiên Chúa đã cho sống lại từ cõi chết.)", italic=True, size=9, space_after=4)
    
    add_p_mixed(doc, [
        ("ĐC: ", True, True),
        ("Lạy Chúa, xin thương xót nghe lời con cầu khẩn.", True, True)
    ], space_after=4)
    
    add_p(doc, "2 Lạy Thiên Chúa là đèn trời soi xét\nKhi con kêu, nguyện Chúa đáp lời.*\nLúc ngặt nghèo, Chúa đã mở lối thoát cho con,\nXin thương xót nghe lời con cầu khẩn.", space_after=4, size=10, line_spacing=1.2)
    add_p(doc, "3 Phàm nhân hỡi, cho đến bao giờ\nLòng vẫn còn chai đá *\nƯa thích chuyện hư không,\nChạy theo điều giả dối?", space_after=4, size=10, line_spacing=1.2)
    add_p(doc, "4 Hãy biết rằng:\nChúa biệt đãi những người trung hiếu với Chúa; *\nKhi tôi kêu, Chúa đã nghe lời.", space_after=4, size=10, line_spacing=1.2)
    add_p(doc, "5 Hãy run sợ và đừng phạm tội nữa,\nTrên giường nằm, hãy suy nghĩ và lặng thinh. *", space_after=4, size=10, line_spacing=1.2)
    add_p(doc, "6 Hãy tiến dâng lễ Chúa như luật truyền\nVà tin tưởng vào Chúa.", space_after=4, size=10, line_spacing=1.2)
    add_p(doc, "7 Biết bao kẻ nói rằng:\n“Ai sẽ cho ta thấy hạnh phúc?”*\nLạy Chúa, xin tỏa ánh tôn nhan của Ngài trên chúng con.", space_after=4, size=10, line_spacing=1.2)
    
    doc.add_page_break() # -> TRANG 9
    
    # =========================================================================
    # TRANG 9: THỨ BẢY - Tv 4 (tiếp) & Tv 133 (134) (Kinh chiều trong đền thánh)
    # =========================================================================
    add_p(doc, "8 Chúa ban xuống lòng con nhiều hoan lạc\nHơn thiên hạ được mùa, lúa rượu đầy dư.", space_before=6, space_after=4, size=10, line_spacing=1.2)
    add_p(doc, "9 Thư thái bình an vừa nằm con đã ngủ,*\nvì chỉ có mình Ngài, lạy Chúa,\n10 Ban cho con được sống yên hàn.", space_after=4, size=10, line_spacing=1.2)
    add_p(doc, "(Vinh Danh Chúa Cha và Chúa Con cùng Vinh Danh Thánh Thần Thiên Chúa. Tự muôn đời và chính hiện nay, luôn mãi đến thiên thu vạn đại. Amen.)", italic=True, size=9, space_after=10)
    
    add_p(doc, "Tv 133 (134) - Kinh chiều trong đền thánh", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=2)
    add_p(doc, "Nào ca ngợi Chúa đi, hỡi tất cả tôi trung của Chúa,\nhỡi những ai lớn nhỏ hằng kính sợ Người (Kh 19,5).", italic=True, space_after=4, size=9)
    
    add_p_mixed(doc, [
        ("ĐC: ", True, True),
        ("Hãy chúc tụng Thiên Chúa\nĐêm đến suốt canh dài.", True, True)
    ], space_after=4)
    
    add_p(doc, "1 Hỡi tất cả những người tôi tớ Chúa\nứng trực suốt đêm trong thánh điện,\nnào chúc tụng Chúa đi !", space_after=4, size=10, line_spacing=1.2)
    add_p(doc, "2 Hãy giơ tay hướng về cung thánh\nMà dâng lên lời chúc tụng Người.", space_after=4, size=10, line_spacing=1.2)
    add_p(doc, "3 Cúi xin Đấng tạo thành trời đất\nXuống cho bạn muôn vàn phúc cả\nTừ núi thánh Xi- on.", space_after=6, size=10, line_spacing=1.2)
    
    add_p(doc, "(Vinh Danh Chúa Cha và Chúa Con cùng Vinh Danh Thánh Thần Thiên Chúa. Tự muôn đời và chính hiện nay, luôn mãi đến thiên thu vạn đại. Amen.)", italic=True, size=9, space_after=4)
    
    doc.add_page_break() # -> TRANG 10
    
    # =========================================================================
    # TRANG 10: CHÚA NHẬT - Tv 90 (91) (Phần 1: Câu 1 - 9)
    # =========================================================================
    add_heading_2(doc, "CHÚA NHẬT", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=1)
    add_p(doc, "(Tức sau kinh chiều II Chúa nhật,\nvà sau kinh chiều II các ngày lễ trọng).", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9, space_after=4)
    
    add_p(doc, "Tv 90 (91) - Nương bóng Chúa toàn năng", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p(doc, "Thầy đã ban cho anh em quyền năng để chà đạp rắn rết, bọ cạp (Lc 10, 19)", italic=True, size=9, space_after=4)
    
    add_p_mixed(doc, [
        ("ĐC: ", True, True),
        ("Nương cánh Chúa, bạn không sợ cảnh hãi hùng đêm vắng.", True, True)
    ], space_after=4)
    
    add_p(doc, "1 Hỡi ai nương tựa Đấng Tối Cao,\nVà núp bóng Đấng quyền năng tuyệt đối,*\n2 Hãy thưa với Chúa rằng:\n“Lạy Thiên Chúa, Ngài là nơi con náu ẩn,\nLà đồn lũy chở che,\nCon tin tưởng vào Ngài.”", space_after=4)
    
    add_p(doc, "3 Chính Chúa gìn giữ bạn\nKhỏi lưới kẻ thù giăng,\nKhỏi tai ương tàn khốc.", space_after=4)
    
    add_p(doc, "4 Chúa phù trì che chở,\nDưới cánh Người, bạn có chỗ ẩn thân:*\nLòng Chúa tín trung là khiên che thuẫn đỡ.", space_after=4)
    
    add_p(doc, "5 Bạn không sợ cảnh hãi hùng đêm vắng\nHay mũi tên bay giữa ban ngày,*\n6 Cả dịch khí hoành hành trong đêm tối,\nCả ôn thần sát hại lúc ban trưa.", space_after=4)
    
    add_p(doc, "7 Dù tả hữu có ngàn người quỵ ngã,\nDù hai bên có chết cả vạn người, *\nRiêng phần bạn tuyệt nhiên không hề hấn.", space_after=4)
    
    add_p(doc, "8 Mở mắt coi, bạn liền thấy rõ\nThế nào là số phận bọn ác nhân.*\n9 Vì bạn có Chúa làm nơi trú ẩn,\nCó Đấng tối cao làm chỗ nương thân.", space_after=4)
    
    doc.add_page_break() # -> TRANG 11
    
    # =========================================================================
    # TRANG 11: CHÚA NHẬT - Tv 90 (91) (Phần 2: Câu 10 - 16)
    # =========================================================================
    add_p(doc, "10 Bạn sẽ không gặp điều ác hại,\nVà tai ương không bén mảng tới nhà,*\n11 Bởi trưng Người truyền cho thiên sứ\nGiữ gìn bạn trên khắp nẻo đường.", space_before=12, space_after=10, size=10.5, line_spacing=1.25)
    
    add_p(doc, "12 Và thiên sứ sẽ tay đỡ tay nâng\nCho bạn khỏi vấp chân vào đá. *\n13 Bạn có thể giẫm lên hùm thiêng rắn độc,\nĐạp nát đầu sư tử khủng long.", space_after=10, size=10.5, line_spacing=1.25)
    
    add_p(doc, "14 Chúa phán:\n“Kẻ gắn bó cùng Ta\nSẽ được ơn giải thoát\nNgười nhận biết danh Ta\nSẽ được sức phù trì;", space_after=10, size=10.5, line_spacing=1.25)
    
    add_p(doc, "15 “Khi kêu đến Ta, Ta liền đáp lại,\nLúc ngặt nghèo, có Ta ở kề bên.\nTa giải cứu và ban nhiều vinh dự,\n16 Cho sống lâu, tuổi thọ dư đầy\nVà hưởng ơn cứu độ Ta ban.”", space_after=12, size=10.5, line_spacing=1.25)
    
    add_p(doc, "(Vinh Danh Chúa Cha và Chúa Con cùng Vinh Danh Thánh Thần Thiên Chúa. Tự muôn đời và chính hiện nay, luôn mãi đến thiên thu vạn đại. Amen.)", italic=True, size=9.5, space_after=4)
    
    doc.add_page_break() # -> TRANG 12
    
    # =========================================================================
    # TRANG 12: THỨ HAI - Tv 85 (86) (Phần 1: Câu 1 - 9)
    # =========================================================================
    add_heading_2(doc, "THỨ HAI", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
    add_p(doc, "Tv 85 (86) - Lời cầu nguyện trong cơn quẫn bách", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p(doc, "Chúc tụng Thiên Chúa, Đấng luôn nâng đỡ ủi an chúng ta trong mọi cơn thử thách (2Cr 1,3.4)", italic=True, size=9, space_after=4)
    
    add_p_mixed(doc, [
        ("ĐC: ", True, True),
        ("Phần Ngài, muôn lạy Chúa, Ngài chậm giận, lại giàu tình thương và lòng thành tín.", True, True)
    ], space_after=4)
    
    add_p(doc, "1 Lạy Chúa, xin lắng tai và đáp lời con,\nVì thân con nghèo hèn túng quẫn.", space_after=4)
    add_p(doc, "2 Xin Chúa bảo toàn sinh mạng con\nBởi vì con trung hiếu. *\nXin cứu độ tôi tớ Ngài đây,\nHằng tin tưởng nơi Ngài.", space_after=4)
    add_p(doc, "3 Chính Ngài là Thiên Chúa của con,\nXin dủ lòng thương con, lạy Chúa: *\nCon kêu con gọi Chúa suốt ngày.", space_after=4)
    add_p(doc, "4 Lạy Chúa, xin làm cho con được vui thỏa,\nVì con nâng tâm hồn lên tới Chúa.", space_after=4)
    add_p(doc, "5 Lạy Chúa, Ngài nhân hậu khoan hồng,\nGiàu tình thương với mọi kẻ kêu xin; *\n6 Lạy Chúa, xin lắng nghe lời con cầu khẩn,\nTiếng con van nài, xin để ý lưu tâm.", space_after=4)
    add_p(doc, "7 Lâm cảnh ngặt nghèo, con kêu lên Chúa,\nVì Chúa vẫn đáp lời. *\n8 Không một thần linh sánh kịp Ngài, lạy Chúa,\nViệc Ngài làm quả thật vô song.", space_after=4)
    add_p(doc, "9 Lạy Chúa, muôn dân chính tay Ngài tạo dựng\nSẽ về phủ phục trước Thánh Nhan\nVà tôn vinh danh Ngài.", space_after=4)
    
    doc.add_page_break() # -> TRANG 13
    
    # =========================================================================
    # TRANG 13: THỨ HAI - Tv 85 (86) (Phần 2: Câu 10 - 17)
    # =========================================================================
    add_p(doc, "10 Vì Ngài thật cao cả\nVà làm nên những việc lạ lùng;*\nChỉ một mình Ngài là Thiên Chúa.", space_before=8, space_after=4)
    add_p(doc, "11 Xin dạy con đường lối Ngài, lạy Chúa,\nCho con vững bước theo chân lý của Ngài. *\nXin Chúa hướng lòng con,\nĐể con biết một niềm kính tôn Danh Thánh.", space_after=4)
    add_p(doc, "12 Con hết lòng cảm tạ,\nLạy Chúa là Thiên Chúa con thờ,\nThánh nhan Ngài, con mãi mãi tôn vinh, *", space_after=4)
    add_p(doc, "13 Vì tình Chúa thương con như trời như biển,\nNgài đã kéo con ra khỏi vực thẳm âm ty.", space_after=4)
    add_p(doc, "14 Lạy Thiên Chúa, phường kiêu ngạo nổi lên chống đối,\nBè lũ hung tàn tìm hại mạng sống con: *\nChúng đâu có kể chi đến Ngài.", space_after=4)
    add_p(doc, "15 Phần Ngài, muôn lạy Chúa,\nNgài là Thiên Chúa nhân hậu từ bi,*\nNgài chậm giận,\nLại giàu tình thương và lòng thành tín.", space_after=4)
    add_p(doc, "16 Xin đoái nhìn và xót thương con,*\nBan sức mạnh của Ngài và xuống ơn cứu độ\nCho tôi tớ Ngài đây, con của nữ tỳ Ngài.", space_after=4)
    add_p(doc, "17 Xin ban cho con một điềm báo phúc,\nĐể bọn thù ghét con trông thấy mà hổ thẹn,*\nVì, lạy Chúa, chính Ngài giúp đỡ ủi an con.", space_after=5)
    
    add_p(doc, "(Vinh Danh Chúa Cha và Chúa Con cùng Vinh Danh Thánh Thần Thiên Chúa. Tự muôn đời và chính hiện nay, luôn mãi đến thiên thu vạn đại. Amen.)", italic=True, size=9, space_after=4)
    
    doc.add_page_break() # -> TRANG 14
    
    # =========================================================================
    # TRANG 14: THỨ BA - Tv 142 (143) (Phần 1: Câu 1 - 6)
    # =========================================================================
    add_heading_2(doc, "THỨ BA", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
    add_p(doc, "Tv 142 (143), 1-11 - Lời cầu xin lúc gặp hiểm nguy", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p(doc, "Con người được nên công chính không phải vì làm những việc lề luật dậy, nhưng vì tin vào Đức Kitô Giêsu (Gl 2,16).", italic=True, size=9, space_after=4)
    
    add_p_mixed(doc, [
        ("ĐC: ", True, True),
        ("Lạy Chúa, xin đừng ẩn mặt đi, vì con vẫn tin cậy nơi Ngài.", True, True)
    ], space_after=4)
    
    add_p(doc, "1 Lạy Chúa, xin nghe lời con khẩn nguyện,*\nLắng nghe con nài van, bởi Ngài thành tín,\n2 Xin chớ đòi tôi tớ ra xét xử,*\nVì trước thánh nhan Ngài\nChẳng có người nào công chính.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "3 Kẻ thù bách hại con,\nChà đạp con dưới đất,\nĐẩy vào chốn tối tăm\nNhư những người đã chết từ bao thuở.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "4 Hơi thở con chỉ còn thoi thóp,\nNghe con tim giá lạnh trong mình,", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "5 Nhớ ngày xưa tháng cũ,\nCon hoài niệm mọi công trình của Chúa,\nVà gẫm suy việc tay Chúa làm nên. *\n6 Hai tay cầu Chúa giơ lên,\nHồn con khát Chúa như miền đất khô.", space_after=6, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 15
    
    # =========================================================================
    # TRANG 15: THỨ BA - Tv 142 (143) (Phần 2: Câu 7 - 11)
    # =========================================================================
    add_p(doc, "7 Xin mau đáp lời con, lạy Chúa,\nHơi thở con nay đã hầu tàn. *\nXin đừng ẩn mặt đi,\nKẻo con hóa ra người thiên cổ.", space_before=8, space_after=8, size=10.5, line_spacing=1.25)
    add_p(doc, "8 Ngay từ buổi sớm mai,\nXin cho con nghiệm thấy tình thương của Chúa,\nVì con vẫn tin cậy nơi Ngài.*\nXin chỉ dạy đường lối phải theo,\nVì con nâng tâm hồn lên cùng Chúa.", space_after=8, size=10.5, line_spacing=1.25)
    add_p(doc, "9 Xin cứu con thoát khỏi địch thù,\nLạy Chúa, bên Ngài con trú ẩn.", space_after=8, size=10.5, line_spacing=1.25)
    add_p(doc, "10 Điều đẹp ý Ngài, xin dạy con thực hiện,\nBởi Ngài là Thiên Chúa của con.*\nXin thần khí tốt lành của Chúa\nDẫn con đi trên miền đất phẳng phiu.", space_after=8, size=10.5, line_spacing=1.25)
    add_p(doc, "11 Lạy Chúa, vì danh dự của Ngài,\nXin cho con được sống.*\nBởi vì Ngài công chính,\nXin cứu con khỏi bước ngặt nghèo.", space_after=10, size=10.5, line_spacing=1.25)
    
    add_p(doc, "(Vinh Danh Chúa Cha và Chúa Con cùng Vinh Danh Thánh Thần Thiên Chúa. Tự muôn đời và chính hiện nay, luôn mãi đến thiên thu vạn đại. Amen.)", italic=True, size=9.5, space_after=4)
    
    doc.add_page_break() # -> TRANG 16
    
    # =========================================================================
    # TRANG 16: THỨ TƯ - Tv 30 (31), 2-6
    # =========================================================================
    add_heading_2(doc, "THỨ TƯ", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
    add_p(doc, "Tv 30 (31), 2-6 - Lời cầu nguyện tin tưởng", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p(doc, "Lạy Cha, xin phó thác hồn con trong tay Cha (Lc 23,46)", italic=True, size=9, space_after=4)
    
    add_p_mixed(doc, [
        ("ĐC: ", True, True),
        ("Xin cứu độ con và cho con ẩn náu bên Ngài.", True, True)
    ], space_after=4)
    
    add_p(doc, "2 Con ẩn náu bên Ngài, lạy Chúa\nXin đừng để con phải tủi nhục bao giờ.\nBởi vì Ngài công chính, xin giải thoát con,\n3 Ghé tai nghe và mau cứu chữa.*\nXin Ngài nên như núi đá cho con trú ẩn,\nNhư thành trì để cứu độ con.", space_after=7, size=10.5, line_spacing=1.25)
    
    add_p(doc, "4 Núi đá và thành lũy bảo vệ con chính là Chúa,\nVì danh dự Ngài, xin dẫn đường chỉ lối cho con.*\n5 Lưới kẻ thù giăng xin gỡ con ra khỏi,\nVì nơi con trú ẩn, chính là Ngài.", space_after=7, size=10.5, line_spacing=1.25)
    
    add_p(doc, "6 Trong tay Ngài, con xin phó thác hồn con,\nNgài đã cứu chuộc con, lạy Chúa Trời thành tín.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "(Vinh Danh Chúa Cha và Chúa Con cùng Vinh Danh Thánh Thần Thiên Chúa. Tự muôn đời và chính hiện nay, luôn mãi đến thiên thu vạn đại. Amen.)", italic=True, size=9.5, space_after=4)
    
    doc.add_page_break() # -> TRANG 17
    
    # =========================================================================
    # TRANG 17: THỨ TƯ - Tv 129 (130) (Tiếng kêu từ vực thẳm)
    # =========================================================================
    add_p(doc, "Tv 129 (130) - Tiếng kêu từ vực thẳm", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=2)
    add_p(doc, "Chính Đức Giêsu sẽ cứu dân Người khỏi tội lỗi (Mt 1,21)", italic=True, size=9, space_after=4)
    
    add_p_mixed(doc, [
        ("ĐC: ", True, True),
        ("Từ vực thẳm, con kêu lên Ngài, lạy Chúa.", True, True)
    ], space_after=4)
    
    add_p(doc, "1 Từ vực thẳm con kêu lên Ngài, lạy Chúa,\n2 Muôn lạy Chúa, xin Ngài nghe tiếng con.*\nDám xin Ngài lắng tai để ý\nNghe lời con tha thiết nguyện cầu.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "3 Ôi lạy Chúa, nếu như Ngài chấp tội,\nNào có ai đứng vững được chăng? *\n4 Nhưng Chúa vẫn rộng lòng tha thứ\nĐể chúng con biết kính sợ Ngài.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "5 Mong đợi Chúa, tôi hết lòng mong đợi,\nCậy trông ở lời Người.*\n6 Hồn tôi trông chờ Chúa,\nHơn lính canh mong đợi hừng đông.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "Hơn lính canh mong đợi hừng đông,\n7 Trông cậy Chúa đi, Ít-ra-en hỡi,*\nBởi Chúa luôn từ ái một niềm,\nƠn cứu chuộc nơi Ngài chan chứa.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "8 Chính Người sẽ cứu chuộc Ít-ra-en\nCho thoát khỏi tội khiên muôn vàn.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "(Vinh Danh Chúa Cha và Chúa Con cùng Vinh Danh Thánh Thần Thiên Chúa. Tự muôn đời và chính hiện nay, luôn mãi đến thiên thu vạn đại. Amen.)", italic=True, size=9.5, space_after=4)
    
    doc.add_page_break() # -> TRANG 18
    
    # =========================================================================
    # TRANG 18: THỨ NĂM - Tv 15 (16) (Chúa là phần gia nghiệp)
    # =========================================================================
    add_heading_2(doc, "THỨ NĂM", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=1)
    add_p(doc, "Tv 15 (16) - Chúa là phần gia nghiệp", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1, size=9.5)
    add_p(doc, "Thiên Chúa đã giải thoát Đức Giêsu khỏi những đau khổ do thần chết gây nên mà cho Người sống lại (Cv 2,24).", italic=True, size=8.5, space_after=3)
    
    add_p_mixed(doc, [
        ("ĐC: ", True, True, False, 9.5),
        ("Thân xác con nghỉ ngơi an toàn.", True, True, False, 9.5)
    ], space_after=3)
    
    add_p(doc, "1 Lạy Chúa Trời, xin giữ gìn con, vì bên Ngài con đang ẩn náu.*\n2 Con thưa cùng Chúa: “Ngài là Chúa con thờ, ngoài Chúa ra đâu là hạnh phúc?”", space_after=2.5, size=9, line_spacing=1.1)
    add_p(doc, "3 Còn thần ngoại xứ này, những thần linh xưa con sùng mộ,*\n4 Vẫn gia tăng tàn phá, và thiên hạ tới tấp tìm theo.\nMáu tế thần, con quyết chẳng dâng, tên của thần, môi con không tụng niệm!", space_after=2.5, size=9, line_spacing=1.1)
    add_p(doc, "5 Lạy Chúa, Chúa là phần sản nghiệp con được hưởng, là chén phúc lộc dành cho con;* Số mạng con chính Ngài nắm giữ.\n6 Phần tuyệt hảo may mắn đã về con, vâng, gia nghiệp ấy làm con thỏa mãn.", space_after=2.5, size=9, line_spacing=1.1)
    add_p(doc, "7 Con chúc tụng Chúa hằng thương chỉ dạy, ngay cả đêm trường lòng dạ nhắn nhủ con.*\n8 Con luôn nhớ có Ngài trước mặt, được Ngài ở bên chẳng nao núng bao giờ.", space_after=2.5, size=9, line_spacing=1.1)
    add_p(doc, "9 Vì thế, tâm hồn con mừng rỡ và lòng dạ hân hoan,* Thân xác con cũng nghỉ ngơi an toàn.\n10 Vì Chúa chẳng đành bỏ mặc con trong cõi âm ty, không để kẻ hiếu trung này hư nát trong phần mộ.", space_after=2.5, size=9, line_spacing=1.1)
    add_p(doc, "11 Chúa sẽ dạy con biết đường về cõi sống:*\nTrước Thánh Nhan, ôi sung sướng tràn trề,\nỞ bên Ngài hoan lạc chẳng hề vơi!", space_after=3, size=9, line_spacing=1.1)
    
    add_p(doc, "(Vinh Danh Chúa Cha và Chúa Con cùng Vinh Danh Thánh Thần Thiên Chúa. Tự muôn đời và chính hiện nay, luôn mãi đến thiên thu vạn đại. Amen.)", italic=True, size=8.5, space_after=2)
    
    doc.add_page_break() # -> TRANG 19
    
    # =========================================================================
    # TRANG 19: THỨ SÁU - Tv 87 (88) (Phần 1: Câu 2 - 10)
    # =========================================================================
    add_heading_2(doc, "THỨ SÁU", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
    add_p(doc, "Tv 87 (88) - Lời cầu xin trong lúc ngặt nghèo", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p(doc, "Đây là giờ các ông hành động, là lúc thần tăm tối hoành hành (Lc 22,53)", italic=True, size=9, space_after=4)
    
    add_p_mixed(doc, [
        ("ĐC: ", True, True),
        ("Trước Thánh Nhan, lạy Chúa, con kêu đêm ngày.", True, True)
    ], space_after=4)
    
    add_p(doc, "2 Lạy Chúa là Thiên Chúa cứu độ con,\nTrước Thánh Nhan, đêm ngày con kêu cứu.*\n3 Nguyện cho lời kinh vọng tới Ngài,\nXin lắng nghe tiếng lòng thổn thức.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "4 Vì hồn con ngập tràn đau khổ,\nMạng sống con âm phủ gần kề,*\n5 Thân thể như đã vào phần mộ,\nVì tựa người kiệt sức còn chi!", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "6 Con nằm đây giữa bao người chết,*\nNhư các tử thi vùi trong mồ mả\nĐã bị Chúa quên đi\nVà không được tay Ngài săn sóc.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "7 Chúa hạ con xuống tận đáy huyệt sâu,\nGiữa chốn tối tăm, giữa lòng huyệt thẳm.*\n8 Cơn giận Chúa đè nặng thân con\nNhư sóng cồn xô đẩy dập vùi.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "9 Chúa làm cho bạn bè xa lánh\nVà coi con như đồ ghê tởm.*\nCon bị giam cầm không thể thoát ra,\n10 Mắt mờ đi vì quá nhiều đau khổ.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "Lạy Chúa, suốt cả ngày con kêu lên Chúa\nVà giơ tay hướng thẳng về Ngài.*", space_after=6, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 20
    
    # =========================================================================
    # TRANG 20: THỨ SÁU - Tv 87 (88) (Phần 2: Câu 11 - 18)
    # =========================================================================
    add_p(doc, "11 Chúa đâu làm phép lạ\nCho người đã mạng vong,\nÂm hồn đâu chỗi dậy\nCa tụng Chúa bao giờ?", space_before=8, space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "12 Trong mồ mả, ai nói về tình thương của Chúa,\nCõi âm ty, ai kể lại lòng thành tín của Ngài?*\n13 Những kỳ công Chúa, nơi tối tăm ai rõ,\nĐức công chính Ngài, chốn quên lãng ai hay?", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "14 Phần con đây, con kêu lên Ngài, lạy Chúa,\nMới tinh sương đã chờ chực nguyện xin.*\n15 Lạy Chúa, thân con đây Chúa nỡ nào ruồng rẫy\nẨn mặt đi mà chẳng đoái hoài.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "16 Từ thuở nhỏ, con khổ đã nhiều và luôn ngắc ngoải,\nChúa làm con kinh hãi, con hóa ra thẫn thờ.\n17 Bao cơn thịnh nộ, Ngài đổ ngập thân con,\nBấy nỗi kinh hoàng khiến con rời rã.", space_after=6, size=10.5, line_spacing=1.25)
    add_p(doc, "18 Bủa vây con suốt ngày ngần ấy thứ,\nDồn dập tư bề như nước bao la.\nCận thân Chúa khiến lìa xa,\nChung quanh bầu bạn chỉ là bóng đêm.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "(Vinh Danh Chúa Cha và Chúa Con cùng Vinh Danh Thánh Thần Thiên Chúa. Tự muôn đời và chính hiện nay, luôn mãi đến thiên thu vạn đại. Amen.)", italic=True, size=9.5, space_after=4)
    
    doc.add_page_break() # -> TRANG 21
    
    # =========================================================================
    # TRANG 21: LỜI CHÚA (Thứ Bảy, Chúa Nhật, Thứ Hai, Thứ Ba)
    # =========================================================================
    add_heading_1(doc, "Lời Chúa", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=4)
    
    add_p(doc, "Thứ bảy Đnl 6, 4- 7", bold=True, italic=True, space_after=3, size=10.5)
    add_p(doc, "Nghe đây hỡi Ít- ra- en, Đức Chúa, Thiên Chúa chúng ta, là Đức Chúa duy nhất. Hãy yêu mến Đức Chúa, Thiên Chúa của anh em, hết lòng hết dạ, hết sức anh em. Những lời này tôi truyền cho anh em hôm nay, anh em phải ghi lòng tạc dạ. Anh em phải lặp lại những lời đó cho con cái, phải nói lại cho chúng, lúc ngồi trong nhà cũng như lúc đi đường, khi đi ngủ cũng như khi thức dậy.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Chúa nhật Kh 22, 4-5", bold=True, italic=True, space_after=3, size=10.5)
    add_p(doc, "Các tôi tớ Thiên Chúa sẽ được nhìn thấy tôn nhan Người, và thánh danh Người ghi trên trán họ. Sẽ không còn đêm tối nữa, họ sẽ không cần đèn đuốc hay ánh sáng mặt trời, vì Đức Chúa là Thiên Chúa sẽ chiếu sáng trên họ, và họ sẽ hiển trị đến muôn thuở muôn đời.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Thứ hai 1 Tx 5, 9- 10", bold=True, italic=True, space_after=3, size=10.5)
    add_p(doc, "Thiên Chúa đã không định cho chúng ta phải chịu cơn thịnh nộ, nhưng được hưởng ơn cứu độ, nhờ Đức Giê- su Ki- tô, Chúa chúng ta, Đấng đã chết vì chúng ta, để dù thức hay ngủ chúng ta cũng cùng sống với Người.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Thứ ba 1 Pr 5,8- 9a", bold=True, italic=True, space_after=3, size=10.5)
    add_p(doc, "Anh em hãy sống tiết độ và tỉnh thức, vì ma quỷ, thù địch của anh em, như sư tử gầm thét, rảo quanh tìm mồi cắn xé. Anh em hãy đứng vững trong đức tin mà chống cự.", space_after=6, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 22
    
    # =========================================================================
    # TRANG 22: LỜI CHÚA (Thứ Tư, Thứ Năm, Thứ Sáu)
    # =========================================================================
    add_p(doc, "Thứ tư Ep 4,26- 27", bold=True, italic=True, space_before=16, space_after=4, size=11)
    add_p(doc, "Anh em đừng phạm tội: chớ để mặt trời lặn mà cơn giận vẫn còn. Đừng để ma quỷ thừa cơ lợi dụng!", space_after=24, size=11, line_spacing=1.3)
    
    add_p(doc, "Thứ năm 1 Tx 5,23", bold=True, italic=True, space_after=4, size=11)
    add_p(doc, "Nguyện chính Thiên Chúa là nguồn mạch bình an, thánh hóa toàn diện con người anh em, để thần trí, tâm hồn và thân xác anh em, được gìn giữ vẹn toàn, không gì đáng trách, trong ngày Đức Giê- su Ki- tô, Chúa chúng ta, quang lâm.", space_after=24, size=11, line_spacing=1.3)
    
    add_p(doc, "Thứ sáu Gr 14,9", bold=True, italic=True, space_after=4, size=11)
    add_p(doc, "Lạy Chúa, Ngài ngự giữa chúng con, và chúng con thuộc về Chúa vì được mang danh Ngài. Xin đừng bỏ rơi chúng con, lạy Chúa là Thiên Chúa chúng con.", space_after=24, size=11, line_spacing=1.3)
    
    add_p(doc, "(sau khi nghe lời Chúa, thinh lặng suy niệm trong giây lát).", italic=True, space_after=6, size=9.5)
    
    doc.add_page_break() # -> TRANG 23
    
    # =========================================================================
    # TRANG 23: XƯỚNG ĐÁP & CA CHÚC TỤNG TIN MỪNG (NUNC DIMITTIS)
    # =========================================================================
    add_heading_1(doc, "Xướng đáp", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=3)
    
    add_p(doc, "Ngoài Mùa Phục Sinh", bold=True, italic=True, space_after=1)
    add_p(doc, "XĐ Trong tay Ngài, lạy Chúa\n* Con xin phó thác hồn con. (đứng lên)\nX Ngài đã cứu chuộc chúng con,\nLạy Chúa Trời thành tín.*", space_after=4)
    
    add_p(doc, "Mùa Phục Sinh", bold=True, italic=True, space_after=1)
    add_p(doc, "XĐ Trong tay Ngài, lạy Chúa,\nCon xin phó thác hồn con.\n* Ha- lê- lui- a. Ha- lê- lui- a.\nX Ngài đã cứu chuộc chúng con,\nLạy Chúa Trời thành tín.*", space_after=6)
    
    add_p(doc, "Tc Tin Mừng “Muôn lạy Chúa”", bold=True, italic=True, space_after=1)
    add_p(doc, "(Làm dấu)", italic=True, size=9, space_after=2)
    
    add_p(doc, "ĐC Lạy Chúa, lúc chúng con còn thức,\nXin Ngài cứu vớt cho,*\nKhi chúng con đã ngủ,\nXin Chúa cũng giữ gìn,*\nĐể cùng thức tỉnh với Đức Ki- tô,\nVà nghỉ ngơi an bình (Ha- lê- lui- a.)", space_after=4)
    
    add_p(doc, "Muôn lạy Chúa, giờ đây\nTheo lời Ngài đã hứa,\nXin để tôi tớ này\nĐược an bình ra đi.", space_after=3)
    
    add_p(doc, "Vì chính mắt con được thấy ơn cứu độ\nChúa đã dành sẵn cho muôn dân: *\nĐó là ánh sáng soi đường cho dân ngoại,\nLà vinh quang của Ít- ra- en dân Ngài.", space_after=3)
    
    add_p(doc, "(Đọc xong làm dấu)", italic=True, size=9, space_after=3)
    
    doc.add_page_break() # -> TRANG 24
    
    # =========================================================================
    # TRANG 24: LỜI NGUYỆN KẾT THÚC (Thứ Bảy, Chúa Nhật, Thứ Hai)
    # =========================================================================
    add_heading_1(doc, "Lời nguyện", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=4, space_after=2)
    add_p(doc, "Chủ sự: Chúng ta dâng lời cầu nguyện", bold=True, italic=True, space_after=4)
    
    add_p(doc, "Thứ bảy", bold=True, italic=True, space_after=2)
    add_p(doc, "Lạy Chúa, xin viếng thăm chúng con đêm nay, để sớm mai, nhờ quyền năng Chúa chúng con được thức dậy và hân hoan mừng ngày Đức Ki tô phục sinh. Người hằng sống và hiển trị muôn đời.", space_after=3)
    
    add_p(doc, "(sau Kinh chiều I các ngày lễ trọng không trùng với ngày Chúa nhật thì đọc)", italic=True, size=9, space_after=2)
    add_p(doc, "Lạy Chúa, xin thăm viếng nhà này, và đuổi xa mọi âm mưu ma quỷ. Xin sai thiên thần Chúa đến ở nơi đây, để gìn giữ chúng con bình an, và xin tuôn đổ phúc lành trên chúng con luôn mãi. Chúng con cầu xin.", space_after=4)
    
    add_p(doc, "Chúa nhật", bold=True, italic=True, space_after=2)
    add_p(doc, "Lạy Chúa, chúng con khiêm tốn dâng lời khẩn nguyện: Hôm nay chúng con đã mừng mầu nhiệm Đức Ki- tô phục sinh, xin cho chúng con thoát mọi điều tai ác, được Chúa cho nghỉ ngơi bình an, và thức dậy hân hoan ca tụng Chúa. Chúng con cầu xin.", space_after=3)
    
    add_p(doc, "(Sau Kinh Chiều II các ngày lễ trọng không trùng ngày Chúa nhật thì đọc)", italic=True, size=9, space_after=2)
    add_p(doc, "Lạy Chúa, xin thăm viếng nhà này, và đuổi xa mọi âm mưu ma quỷ. Xin sai thiên thần Chúa đến ở nơi đây, để gìn giữ chúng con bình an, và xin tuôn đổ phúc lành trên chúng con luôn mãi. Chúng con cầu xin.", space_after=4)
    
    add_p(doc, "Thứ hai", bold=True, italic=True, space_after=2)
    add_p(doc, "Lạy Chúa ngày hôm nay chúng con đã khó nhọc gieo hạt giống Nước Trời, giờ đây xin cho chúng con được nghỉ ngơi lại sức, và xin cho hạt giống chúng con gieo vãi được nảy nở và trổ bông chín vàng trong ngày mùa sau hết. Chúng con cầu xin.", space_after=3)
    
    doc.add_page_break() # -> TRANG 25
    
    # =========================================================================
    # TRANG 25: LỜI NGUYỆN KẾT THÚC (Thứ Ba, Thứ Tư, Thứ Năm, Thứ Sáu)
    # =========================================================================
    add_p(doc, "Thứ ba", bold=True, italic=True, space_before=8, space_after=3, size=11)
    add_p(doc, "Lạy Chúa, xin thương chiếu giải ánh sáng Chúa vào bóng tối đêm nay, và ban cho chúng con là con cái Chúa được nghỉ ngơi an lành, để nhân danh Chúa, chúng con được vui mừng thức dậy hưởng ánh quang ngày mới. Chúng con cầu xin.", space_after=10, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Thứ tư", bold=True, italic=True, space_after=3, size=11)
    add_p(doc, "Lạy Chúa Giê su Ki- tô khiêm nhường và hiền hậu, ách của Chúa êm ái, gánh của Chúa nhẹ nhàng. Chúng con đến trao cho Chúa gánh nặng của ngày hôm nay, xin cho chúng con được nghỉ ngơi bên Chúa là Đấng hằng sống và hiển trị muôn đời.", space_after=10, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Thứ năm", bold=True, italic=True, space_after=3, size=11)
    add_p(doc, "Lạy Chúa là Thiên Chúa chúng con, xin cho chúng con qua một đêm yên giấc, và xin dùng ánh sáng của Chúa thức tỉnh chúng con, để chúng con hân hoan nhìn thấy ngày mới rạng rỡ huy hoàng, nhân danh Đức Ki tô Chúa chúng con. Amen.", space_after=10, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Thứ sáu", bold=True, italic=True, space_after=3, size=11)
    add_p(doc, "Lạy Chúa nhân từ, Chúa đã dùng cái chết của Con Một Chúa mà giao hòa thế gian với Chúa. Giờ đây, xin đoái thương gìn giữ chúng con, cho chúng con được ngủ yên hàn, để mai ngày thức dậy, chúng con vui mừng ca tụng vinh quang Chúa, Đấng cứu độ chúng con. Người hằng sống và hiển trị muôn đời. Amen.", space_after=6, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 26
    
    # =========================================================================
    # TRANG 26: LẦN CHUỖI KÍNH ĐỨC MẸ (Năm Sự Vui & Năm Sự Sáng)
    # =========================================================================
    add_heading_1(doc, "LẦN CHUỖI KÍNH ĐỨC MẸ", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=3)
    
    add_p(doc, "Năm sự Vui", bold=True, italic=True, space_after=3, size=10.5)
    add_p(doc, "Thứ nhất thì ngắm: Thiên Thần truyền tin cho Đức Bà chịu thai. Ta hãy xin cho được ở khiêm nhường.\nThứ hai thì ngắm: Đức Bà đi viếng Bà thánh Isave. Ta hãy xin cho được lòng yêu người.\nThứ ba thì ngắm: Đức Bà sinh Đức Chúa Giêsu trong hang đá. Ta hãy xin cho được lòng khó khăn.\nThứ bốn thì ngắm: Đức Bà dâng Đức Chúa Giêsu trong đền thánh. Ta hãy xin cho được vâng lời chịu lụy.\nThứ năm thì ngắm: Đức Bà tìm được Đức Chúa Giêsu trong đền thánh. Ta hãy xin cho được giữ nghĩa cùng Chúa luôn.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Năm sự Sáng", bold=True, italic=True, space_after=3, size=10.5)
    add_p(doc, "Thứ nhất thì ngắm: Đức Chúa Giêsu chịu phép Rửa tại sông Gio-đan. Ta hãy xin cho được sống xứng đáng là con cái Chúa.\nThứ hai thì ngắm: Đức Chúa Giêsu dự tiệc cưới Cana. Ta hãy xin cho được noi gương Đức Mẹ mà vững tin vào Chúa.\nThứ ba thì ngắm: Đức Chúa Giêsu rao giảng Nước Trời và kêu gọi sám hối. Ta hãy xin cho được tin vào Tin Mừng và biết hoán cải đời sống.\nThứ bốn thì ngắm: Đức Chúa Giêsu biến hình trên núi. Ta hãy xin cho được biến đổi nhờ lắng nghe Lời Chúa.\nThứ năm thì ngắm: Đức Chúa Giêsu lập Bí Tích Thánh Thể. Ta hãy xin cho được sốt sắng tham dự thánh lễ và rước Mình Thánh Chúa.", space_after=4, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 27
    
    # =========================================================================
    # TRANG 27: LẦN CHUỖI KÍNH ĐỨC MẸ (Năm Sự Thương & Năm Sự Mừng)
    # =========================================================================
    add_p(doc, "Năm sự Thương", bold=True, italic=True, space_before=6, space_after=3, size=10.5)
    add_p(doc, "Thứ nhất thì ngắm: Đức Chúa Giêsu lo buồn đổ mồ hôi máu. Ta hãy xin cho được ăn năn tội nên.\nThứ hai thì ngắm: Đức Chúa Giêsu chịu đánh đòn. Ta hãy xin cho được hãm mình chịu khó bằng lòng.\nThứ ba thì ngắm: Đức Chúa Giêsu chịu đội mão gai. Ta hãy xin cho được chịu mọi sự sỉ nhục bằng lòng.\nThứ bốn thì ngắm: Đức Chúa Giêsu vác Cây Thánh Giá. Ta hãy xin cho được vác Thánh Giá theo chân Chúa.\nThứ năm thì ngắm: Đức Chúa Giêsu chịu chết trên Cây Thánh Giá. Ta hãy xin đóng đinh tính xác thịt vào Thánh Giá Chúa.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "Năm sự Mừng", bold=True, italic=True, space_after=3, size=10.5)
    add_p(doc, "Thứ nhất thì ngắm: Đức Chúa Giêsu sống lại. Ta hãy xin cho được sống lại thật về phần linh hồn.\nThứ hai thì ngắm: Đức Chúa Giêsu lên trời. Ta hãy xin cho được ái mộ những sự trên trời.\nThứ ba thì ngắm: Đức Chúa Thánh Thần hiện xuống. Ta hãy xin cho được lòng đầy dẫy mọi ơn Đức Chúa Thánh Thần.\nThứ bốn thì ngắm: Đức Chúa Trời cho Đức Bà lên trời. Ta hãy xin ơn chết lành trong tay Đức Mẹ.\nThứ năm thì ngắm: Đức Chúa Trời thưởng Đức Mẹ trên trời. Ta hãy xin Đức Mẹ phù hộ cho ta được thưởng cùng Đức Mẹ trên nước Thiên đàng.", space_after=4, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 28
    
    # =========================================================================
    # TRANG 28: KINH LẠY NỮ VƯƠNG & BÀI HÁT KÍNH ĐỨC MẸ 1 (Năm xưa trên cây sồi)
    # =========================================================================
    add_p(doc, "Kinh Lạy Nữ Vương", bold=True, italic=True, space_before=4, space_after=3, size=11)
    add_p(doc, "Lạy Nữ Vương Mẹ nhân lành làm cho chúng con được sống, được vui, được cậy. Thân lạy Mẹ, chúng con, con cháu E-và ở chốn khách đầy, kêu đến cùng Bà; Chúng con ở nơi khóc lóc than thở kêu khấn Bà thương. Hỡi ôi! Bà là Chúa bầu chúng con, xin ghé mặt thương xem chúng con. Đến sau khỏi đày, xin cho chúng con được thấy Đức Chúa Giêsu, Con lòng Bà gồm phúc lạ. Ôi khoan thay! nhân thay! dịu thay! Thánh Maria trọn đời đồng trinh. Amen.", space_after=12, size=10.5, line_spacing=1.25)
    
    add_heading_1(doc, "BÀI HÁT KÍNH ĐỨC MẸ", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=4)
    add_p(doc, "1. Năm xưa trên cây sồi", bold=True, italic=True, space_after=3, size=11)
    add_p(doc, "1. Năm xưa trên cây sồi, làng Fatima xa xôi. Có Đức Mẹ Chúa Trời, hiện ra uy linh sáng chói. Mẹ nhắn nhủ người đời: “Hãy mau ăn năn đền bồi, hãy tôn sùng Mẫu Tâm, hãy năng lần hạt Mân Côi.”", space_after=5, size=10, line_spacing=1.2)
    add_p_mixed(doc, [
        ("ĐK. ", True, False),
        ("Mẹ Maria ơi, Mẹ Maria ơi! Con vâng nghe Mẹ rồi. Sớm chiều từ nay thống hối. Mẹ Maria ơi! Xin Mẹ đoái thương nhận lời, cho Nước Việt xinh tươi, đức tin sáng ngời.", False, False)
    ], space_after=5, line_spacing=1.2)
    add_p(doc, "2. Đôi môi như hoa cười, Mẹ Maria vui tươi. Có biết bao lớp người, gần xa đua nhau bước tới. Lòng trút khỏi ngậm ngùi, mắt khô đôi suối lệ đời. Ngước trông về Mẫu Tâm, sống trong tình Mẹ yên vui.", space_after=4, size=10, line_spacing=1.2)
    
    doc.add_page_break() # -> TRANG 29
    
    # =========================================================================
    # TRANG 29: BÀI HÁT KÍNH ĐỨC MẸ 2 (Hoa Mân Côi - trọn vẹn cả bài)
    # =========================================================================
    add_p(doc, "2. Hoa Mân Côi", bold=True, italic=True, space_before=16, space_after=8, size=12)
    add_p_mixed(doc, [
        ("ĐK. ", True, False, False, 11),
        ("Một tràng hoa Mân Côi, hương thơm ngào ngát thắm ân tình. Từng lời kinh sốt mến, xin dâng về Nữ Vương hiển vinh.", False, False, False, 11)
    ], space_after=12, line_spacing=1.28)
    add_p(doc, "1. Khi an vui thái hoà, như ngày Thiên sứ truyền tin. Như trong đêm Ngôi Lời giáng trần, như khi Mẹ gặp lại Con Chúa. Kết hoa kinh để dâng Mẹ.", space_after=12, size=11, line_spacing=1.28)
    add_p(doc, "2. Khi đau thương lấp đầy, như vườn loang máu mồ hôi. Như gai đâm, như chịu nát đòn, như khi nhận thập tự hy tế. Kết hoa kinh để dâng Mẹ.", space_after=12, size=11, line_spacing=1.28)
    add_p(doc, "3. Khi hân hoan cõi lòng, như ngày Con Chúa phục sinh. Như khi môn sinh nhận Thánh thần, như khi Mẹ về trời vinh sáng. Kết hoa kinh để dâng Mẹ.", space_after=8, size=11, line_spacing=1.28)
    
    doc.add_page_break() # -> TRANG 30
    
    # =========================================================================
    # TRANG 30: BÀI HÁT KÍNH ĐỨC MẸ 3 (Lạy Mẹ Maria, Mẹ Thiên Chúa)
    # =========================================================================
    add_p(doc, "3. Lạy Mẹ Maria, Mẹ Thiên Chúa", bold=True, italic=True, space_before=16, space_after=8, size=12)
    add_p_mixed(doc, [
        ("ĐK: ", True, False, False, 11),
        ("Lạy Mẹ Maria Mẹ Thiên Chúa Mẹ đồng trinh. Đoàn con chung tiếng hát dâng tấm lòng, dâng đời sống. Lạy Mẹ Maria Mẹ nhân ái, Mẹ hiển vinh. Mẹ chính là Nữ Vương, là Trạng Sư, là Mẹ con.", False, False, False, 11)
    ], space_after=14, line_spacing=1.28)
    add_p(doc, "1. Con dâng Mẹ đây tâm hồn, đây trí khôn, cả dĩ vãng cả hiện tại với tương lai. Đức thiện toàn con cương quyết gắng đi tới. Trông lên Mẹ là gương mẫu của đời con.", space_after=14, size=11, line_spacing=1.28)
    add_p(doc, "2. Yêu thanh bần, yêu vâng lời, yêu khiết trinh, và yêu sống trên con đường Chúa đi xưa. Xứng con Mẹ con vui bước tới Thiên Chúa. Hy sinh nhiều vì bác ái quên lợi danh.", space_after=8, size=11, line_spacing=1.28)
    
    doc.add_page_break() # -> TRANG 31
    
    # =========================================================================
    # TRANG 31: BÀI HÁT KÍNH ĐỨC MẸ 4 (Xin Vâng - trọn vẹn cả bài)
    # =========================================================================
    add_p(doc, "4. Xin Vâng", bold=True, italic=True, space_before=16, space_after=8, size=12)
    add_p(doc, "1. Mẹ ơi đời con dõi bước theo Mẹ, lòng con quyết noi gương Mẹ, xin Mẹ dạy con hai tiếng: ‘Xin Vâng’. Mẹ ơi đường đi trăm ngàn nguy khó, hiểm nguy dâng tràn đây đó, xin Mẹ dạy con hai tiếng: ‘Xin Vâng’.", space_after=12, size=11, line_spacing=1.28)
    add_p_mixed(doc, [
        ("ĐK: ", True, False, False, 11),
        ("Xin vâng, Mẹ dạy con hai tiếng ‘Xin Vâng’, hôm qua, hôm nay và ngày mai. Xin vâng, Mẹ dạy con hai tiếng ‘Xin Vâng’, hôm nay, tương lai và suốt đời.", False, False, False, 11)
    ], space_after=12, line_spacing=1.28)
    add_p(doc, "2. Mẹ ơi đời con dõi bước theo Mẹ, lòng con quyết noi gương Mẹ, xin Mẹ dạy con hai tiếng: ‘Xin Vâng’. Mẹ ơi đường đi phủ đầy bóng tối, bẫy chông giăng tràn muôn lối, xin Mẹ dạy con hai tiếng: ‘Xin Vâng’.", space_after=8, size=11, line_spacing=1.28)
    
    doc.add_page_break() # -> TRANG 32
    
    # =========================================================================
    # TRANG 32: BÀI HÁT KÍNH ĐỨC MẸ 5 (Ngày Nay Con Đến)
    # =========================================================================
    add_p(doc, "5. Ngày Nay Con Đến", bold=True, italic=True, space_before=16, space_after=8, size=12)
    add_p_mixed(doc, [
        ("ĐK: ", True, False, False, 11),
        ("Ngày nay con đến hát khen mừng Mẹ Chúa thiên đàng, dâng ngành Mân Côi muôn mầu hoa thắm tươi. Lạy Mẹ yêu mến, lắng nghe lời con hát nhịp nhàng, hòa với cung đàn xiết bao mừng vui. Ôi Maria, phúc đức no đầy chan hòa, lòng con yêu mến cậy trông thiết tha, qua cơn gian nan, giữa chốn sa trường nguy biến, xin đưa hồn con tới quê thanh nhàn.", False, False, False, 11)
    ], space_after=14, line_spacing=1.28)
    add_p(doc, "1. Mẹ ơi, lời Mẹ thiết tha nài xin, con năng ngắm phép Mân Côi từ đây. Này con thành tâm mến yêu cậy tin, cao rao phép thánh Mân Côi hằng ngày.", space_after=14, size=11, line_spacing=1.28)
    add_p(doc, "2. Mẹ ơi trần gian biết bao lầm than, tan theo năm tháng chiến tranh còn chi ? Này con cậy trông Nữ Vương bình an, xin ơn phép thánh Mân Côi phù trì.", space_after=8, size=11, line_spacing=1.28)
    
    doc.add_page_break() # -> TRANG 33
    
    # =========================================================================
    # TRANG 33: BÀI HÁT KÍNH ĐỨC MẸ 6 (Tận Hiến Cho Mẹ - trọn vẹn cả bài)
    # =========================================================================
    add_p(doc, "6. Tận Hiến Cho Mẹ", bold=True, italic=True, space_before=16, space_after=8, size=12)
    add_p(doc, "1. Con đến trước tòa Nữ Vương uy quyền, dâng hồn dâng xác, dâng cõi lòng yêu mến, phó trót nơi Mẹ tấm thân nhỏ hèn, để đời con luôn vui sống bằng yên.", space_after=12, size=11, line_spacing=1.28)
    add_p_mixed(doc, [
        ("ĐK: ", True, False, False, 11),
        ("Ôi Maria xin Mẹ nhận lấy, tấm thân khổ hèn con đến kính dâng , quyết chí thánh hóa nhờ Mẹ với Mẹ, vững chí chiến đấu vì Mẹ trong Mẹ, Nước Mẹ thống trị chiến sĩ trên đường mới, xây đắp vinh quang Nước Cha muôn đời.", False, False, False, 11)
    ], space_after=12, line_spacing=1.28)
    add_p(doc, "2. Con khấn xin Mẹ những khi đau buồn, trên đường con đi trong những ngày nguy khốn, đôi mắt đau buồn ngước trông lên Mẹ, để được ủi an dưới bóng Mẹ yêu.", space_after=8, size=11, line_spacing=1.28)
    
    doc.add_page_break() # -> TRANG 34
    
    # =========================================================================
    # TRANG 34: BÀI HÁT KÍNH ĐỨC MẸ 7 (Lạy Mẹ Fatima) & 8 (Sống Gần Mẹ)
    # =========================================================================
    add_p(doc, "7. LẠY MẸ FATIMA", bold=True, space_before=6, space_after=3, size=11)
    add_p_mixed(doc, [
        ("ĐK: ", True, False),
        ("Lạy Mẹ Fatima, Mẹ nỉ non bao lần. Tội gian trần để phiền cho trái tim Mẹ. Lòng Mẹ thương bao la. Tình ái ân vô ngần. Con dâng mình đền thay tội lỗi muôn dân.", False, False)
    ], space_after=4, line_spacing=1.25)
    add_p(doc, "1. Từ nay lòng con nhớ lời Mẹ tha thiết, xá gì dầu nguy biến thề chết có Mẹ, vui có Mẹ. Dù rằng bao sóng gió con lo gì, con lo gì. Chết bên Mẹ con sợ chi con sợ gì, Mẹ ơi.", space_after=4, size=10.5, line_spacing=1.25)
    add_p(doc, "2. Mẹ cho lòng con trên đường dài dương thế, tấm lòng luôn trinh khiết tựa hoa trắng ngần, hoa trắng ngần. Trọn đời con trung tín sống với Mẹ, luôn cho Mẹ, biết hy sinh quên lợi danh, sống cuộc đời bình an.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "8. Sống Gần Mẹ", bold=True, italic=True, space_after=3, size=11)
    add_p_mixed(doc, [
        ("ĐK: ", True, False),
        ("Sống gần Mẹ lòng con hoan lạc biết bao, Mẹ ơi ! Sống gần Mẹ lòng con êm đềm thiết tha Mẹ ơi.", False, False)
    ], space_after=4, line_spacing=1.25)
    add_p(doc, "1. Lòng Mẹ ấm êm dịu dàng con muốn đêm ngày đến nương tựa bên lòng. Mẹ con thở than thiết tha, nỗi vui hoặc nỗi buồn, lòng con được thảnh thơi. Đời con chứa chan bao tình âu yếm khi được sống bên Mẹ đêm ngày. Lòng con yên hàn vững tin tới khi về quê trời là cõi vui đời đời.", space_after=6, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 35
    
    # =========================================================================
    # TRANG 35: BÀI HÁT KÍNH ĐỨC MẸ 9 (Tiếng Hát Thiên Thu) & 10 (Thành Tâm Dâng Bài Ca)
    # =========================================================================
    add_p(doc, "9. Tiếng Hát Thiên Thu", bold=True, italic=True, space_before=6, space_after=3, size=11)
    add_p_mixed(doc, [
        ("ĐK: ", True, False),
        ("Trọn đời con một bài hát kính ca ngợi. Tiếng hát con vang tận tới thiên thu. Mẹ tình thương từ trời ngút cao vời vợi. Xin yêu thương đón nhận cả hồn thơ. Xin yêu thương dắt về chốn quê mộng mơ.", False, False)
    ], space_after=4, line_spacing=1.25)
    add_p(doc, "1. Nhìn trời cao thăm thẳm mơ ngày về, xin đoái nhìn chở che năm tháng đời con. Ngước trông trời con tìm theo dáng Mẹ hiền, để mơ ngày về quê sống vui bên Mẹ.", space_after=4, size=10.5, line_spacing=1.25)
    add_p(doc, "2. Mẹ là sao soi nẻo trên dương trần. Đêm tối trời nhìn sao con bước bình an. Nếu khi nào con lạc chân bước sai đường, Mẹ thương tình dìu con bước theo chân Mẹ.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "10. Thành Tâm Dâng Bài Ca Yêu Mến", bold=True, italic=True, space_after=3, size=11)
    add_p_mixed(doc, [
        ("ĐK: ", True, False),
        ("Thành tâm dâng bài ca yêu mến lên Mẹ trên trời, Mẹ ơi cho hồn con vui sống bên Mẹ muôn đời.", False, False)
    ], space_after=4, line_spacing=1.25)
    add_p(doc, "1. Khúc ca dâng tiến Mẹ ngát bay tựa hương trầm, lời ca còn vang vang, hòa ý thơ trìu mến.", space_after=4, size=10.5, line_spacing=1.25)
    add_p(doc, "2. Tiếng con vang tới trời, tán dương Mẹ muôn đời, hòa tâm tình nơi nơi, nguyện lắng nghe Mẹ hỡi.", space_after=6, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 36
    
    # =========================================================================
    # TRANG 36: BÀI HÁT KÍNH ĐỨC MẸ 11 (Con Xin Dâng Mẹ) & 12 (Dâng Mẹ)
    # =========================================================================
    add_p(doc, "11. Con Xin Dâng Mẹ", bold=True, italic=True, space_before=6, space_after=3, size=11)
    add_p_mixed(doc, [
        ("ĐK. ", True, False),
        ("Mẹ ơi trước nhan Mẹ con dâng về Mẹ một tràng hoa Mân Côi, và ngàn lời ca chan chứa tình yêu. Lời con tiếng ca hòa dâng lên Mẹ hiền, tựa ngàn hoa thắm tươi xin dâng lên Mẹ thương yêu.", False, False)
    ], space_after=4, line_spacing=1.25)
    add_p(doc, "1. Trong khi an vui con dâng lên Mẹ tình yêu, xin dâng trọn niềm trìu mến. Khi con cô đơn xin dâng về Mẹ đời những gian truân Mẹ sẽ ủi an.", space_after=4, size=10.5, line_spacing=1.25)
    add_p(doc, "2. Trong khi âu lo xin dâng lên Mẹ lời yêu, con dâng trong niềm phó thác. Khi con bơ vơ trên nơi gian trần nhìn ánh sao mai kêu khấn Mẹ yêu.", space_after=12, size=10.5, line_spacing=1.25)
    
    add_p(doc, "12. Dâng Mẹ", bold=True, italic=True, space_after=4, size=11)
    add_p(doc, "1. Mẹ ơi, con biết lòng Mẹ từ nhân, Mẹ lắng nghe con cầu khẩn. Vì xưa nay, chưa từng có ai kêu cầu, mà Mẹ ngoảnh mặt làm ngơ.", space_after=4, size=10.5, line_spacing=1.25)
    add_p_mixed(doc, [
        ("ĐK. ", True, False),
        ("Mẹ Mẹ ơi con dâng lên Mẹ, gia đình con và giáo xứ con. Mẹ cầu Chúa xuống ơn lành, cho mọi người vui sống an bình. Mẹ Mẹ ơi con dâng lên Mẹ, gia đình con và giáo xứ con. Mẹ cầu Chúa xuống ơn lành cho mọi người sống trong an bình.", False, False)
    ], space_after=4, line_spacing=1.25)
    add_p(doc, "2. Tình con năm tháng một niềm cậy trông, cầu Chúa thương ban nguồn sống. Mẹ thông ơn muôn đời lắng nghe dân Người. Lộc trời sáng ngời ngàn nơi.", space_after=4, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 37
    
    # =========================================================================
    # TRANG 37: BÀI HÁT KÍNH ĐỨC MẸ 13 (Hoa Lòng Dâng Mẹ) & 14 (Mẹ Đẹp Tươi)
    # =========================================================================
    add_p(doc, "13. Hoa Lòng Dâng Mẹ", bold=True, italic=True, space_before=6, space_after=3, size=11)
    add_p(doc, "1. Con đến dâng Mẹ đoá hoa lòng này lạy Mẹ đoàn con dâng tiến. Xin Mẹ, xin Mẹ thánh hiến, ôi lạy Mẹ ấp ủ con liên.", space_after=4, size=10.5, line_spacing=1.25)
    add_p_mixed(doc, [
        ("ĐK. ", True, False),
        ("Ma-ri-a Mẹ tuyệt vời, ôi Mẹ đẹp ngời, con chào Mẹ là Nữ Trinh Vương. Ma-ri-a Mẹ Thiên Chúa, ôi Mẹ con người, xin dâng Mẹ hồn xác con đây.", False, False)
    ], space_after=4, line_spacing=1.25)
    add_p(doc, "2. Lạy Nữ Vương Vô Nhiễm Nguyên Tội, Mẹ dịu dàng từ bi lân ái. Xin Mẹ, xin Mẹ thương đoái xuống dạt dào ơn thánh tuôn trào.", space_after=12, size=10.5, line_spacing=1.25)
    
    add_p(doc, "14. Mẹ Đẹp Tươi", bold=True, italic=True, space_after=4, size=11)
    add_p(doc, "1. Mẹ Ma-ri-a đẹp tươi như bình minh chiếu rạng ngời. Vầng trăng lung linh trời đêm so với Mẹ còn kém xa.", space_after=4, size=10.5, line_spacing=1.25)
    add_p_mixed(doc, [
        ("ĐK. ", True, False),
        ("Con thành tâm kính mừng Mẹ, mừng Mẹ đầy ơn phúc, lòng khiết trinh như huệ thắm xinh muôn đời hiển vinh.", False, False)
    ], space_after=4, line_spacing=1.25)
    add_p(doc, "2. Mẹ như hoa thơm ngàn hương luôn toả bay chín tầng trời, mọi ô nhơ không hề vương, ôi xác hồn Mẹ sáng tươi.", space_after=4, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 38
    
    # =========================================================================
    # TRANG 38: BÀI HÁT KÍNH ĐỨC MẸ 15 (Mẹ Nhân Loại) & 16 (Mẹ Đứng Đó)
    # =========================================================================
    add_p(doc, "15. Mẹ Nhân Loại", bold=True, italic=True, space_before=6, space_after=3, size=11)
    add_p_mixed(doc, [
        ("ĐK. ", True, False),
        ("Ma-ri-a ngày xa xưa ấy, đồi hoang liêu bóng mây nhạt chiều, đứng tiêu điều, xót xa nhiều, Thánh Giá chiều treo xác con yêu.", False, False)
    ], space_after=4, line_spacing=1.25)
    add_p(doc, "1. Mẹ đứng nhìn nhân loại, ôi con người đang hấp hối giữa tội nhơ. Mẹ nhìn lên bàn tay Thánh Giá, chung tâm hồn, chung đau buồn, bên con mình cứu rỗi cho cuộc đời.", space_after=4, size=10.5, line_spacing=1.25)
    add_p(doc, "2. Mẹ đã nhiều mong đợi cho con người mau xa thoát những khổ đau. Mẹ đồng công cùng con thương mến, xin ân tình, xin an bình, như mưa hồng rơi xuống cho lòng người.", space_after=8, size=10.5, line_spacing=1.25)
    
    add_p(doc, "16. MẸ ĐỨNG ĐÓ", bold=True, space_after=3, size=11)
    add_p_mixed(doc, [
        ("ĐK. ", True, False),
        ("Mẹ đứng đó khi hoàng hôn tím màu. Nhạc thương trầm buông hắt hiu đồi cao u hoài loang máu đào. Con Chúa đau thương treo trên thập giá, hiến thân vì nhân loại tội tình. Mẹ đứng đó tâm hồn tê tái sầu, đồng công cùng Con dấu yêu, vì thương nhân loại bao khốn cùng. Xin dẫn đưa bao tâm hồn lạc hướng về bên Mẹ, Mẹ ơi.", False, False)
    ], space_after=4, line_spacing=1.25)
    add_p(doc, "1. Hỡi ai qua đường ngừng bước đây mà chiêm ngắm: Chúa chí tôn cam chịu muôn nỗi khổ đau. Vì thương toàn nhân loại muôn chốn muôn đời. Ngày đêm khóc than mong chờ vinh phúc quê trời.", space_after=4, size=10.5, line_spacing=1.25)
    
    doc.add_page_break() # -> TRANG 39
    
    # =========================================================================
    # TRANG 39: BÀI HÁT KÍNH ĐỨC MẸ 17 (Ngàn Hoa) & KÍNH THÁNH PHANXICÔ
    # =========================================================================
    add_p(doc, "17. Ngàn Hoa", bold=True, italic=True, space_before=4, space_after=2)
    add_p_mixed(doc, [
        ("ĐK. ", True, False),
        ("Ngàn hoa đẹp tươi con dâng lên Nữ Vương. Con dâng lên Nữ Vương cùng tiếng ca nguyện cầu tháng ngày vọng ngân, ôi Mẹ Chúa thiên đường.", False, False)
    ], space_after=2.5)
    add_p(doc, "1. Mẹ ban cho hồn con xinh đẹp tựa hoa tinh trắng, nhuần thắm hương thơm nồng ngát bay nhẹ nhàng tới quê thiên đàng.", space_after=2.5)
    add_p(doc, "2. Đời con nơi trần gian muôn ngàn khổ đau nguy biến. Chỉ biết trông lên Mẹ khấn xin cho hồn lắng vơi ưu phiền.", space_after=5)
    
    add_heading_1(doc, "BÀI HÁT KÍNH THÁNH PHANXICÔ", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=3)
    
    add_p(doc, "Lạy Thánh Phanxicô", bold=True, italic=True, space_after=1.5)
    add_p(doc, "Lạy thánh Phanxicô /người như ánh sao ngời giữa ngàn mây u tối /lạy thánh Phanxicô/ như trăng rằm chỉ lối trong đêm đen /như mặt trời tỏa ánh dịu hiền. Xin ban cho con noi theo bước người đi /cho lòng con không còn mong ước chi / bằng đơn sơ sống đời mến yêu /bằng đơn sơ sống đời mến yêu.", italic=True, space_after=4)
    
    add_p(doc, "Xin thu hút lòng con", bold=True, italic=True, space_after=1.5)
    add_p(doc, "Lạy Chúa xin thu hút lòng con /khỏi đam mê mọi sự trên đời /Xin thu hút lòng con thu hút bằng tình yêu /nóng bỏng và ngọt ngào của Chúa. Để con chết vì say mê mối tình của Ngài //như chính Ngài đã vui lòng chết vì say mê mối tình của con.", italic=True, space_after=4)
    
    add_p(doc, "Lạy Thiên Chúa chí tôn hiển vinh", bold=True, italic=True, space_after=1.5)
    add_p(doc, "Lạy Thiên Chúa, Thiên Chúa chí tôn hiển vinh/ Xin chiếu sáng vào tâm hồn mù tối của con /Xin ban cho con một đức tin chân chính /Cho con một đức cậy vững chắc /Cho con một đức mến hoàn hảo. Xin làm cho con nên minh mẫn tinh tường /Để cho con biết tuân hành thánh ý đích thực của Chúa.", italic=True, space_after=3)
    
    doc.add_page_break() # -> TRANG 40
    
    # =========================================================================
    # TRANG 40: KINH CẦU CHO CÁC ÂN NHÂN VÀ THÂN NHÂN & KẾT THÚC
    # =========================================================================
    add_heading_1(doc, "KINH CẦU NGUYỆN CHO CÁC ÂN NHÂN VÀ THÂN NHÂN", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=12)
    
    add_p_mixed(doc, [
        ("Chủ sự: ", True, True, False, 11),
        ("chúng ta hãy cầu nguyện cho ông bà cha mẹ, anh chị em, thân nhân và ân nhân của chúng ta, cùng tất cả những ai xin chúng ta cầu nguyện, cho kẻ đang sống cũng như kẻ đã qua đời ", False, False, False, 11),
        ("(thinh lặng ít phút)", False, True, False, 11)
    ], space_after=10, line_spacing=1.3)
    
    add_p_mixed(doc, [
        ("Tất cả: ", True, False, False, 11),
        ("Lạy Chúa, với ơn Thánh Linh, Chúa đã đổ lửa mến trong tâm hồn các tín hữu, vì lòng lân tuất Chúa, xin ban cho ông bà cha mẹ, anh chị em, thân nhân và ân nhân chúng con sức mạnh hồn xác để họ được tận lực kính mến Chúa và hết lòng thực hành mọi điều làm đẹp lòng Chúa.", False, False, False, 11)
    ], space_after=10, line_spacing=1.3)
    
    add_p(doc, "Lạy Chúa là Đấng rộng lượng tha thứ và muốn cứu rỗi mọi người. Chúng con xin Chúa khoan hồng cho ông bà cha mẹ, anh chị em, thân nhân, ân nhân của chúng con đã từ biệt cõi đời; xin Chúa vì lời Đức Maria Đồng Trinh và các thánh cầu bầu cho các linh hồn ấy được hưởng vinh phúc vô cùng. Nhờ Đức Giêsu Kitô Chúa chúng con. Amen.", space_after=24, size=11, line_spacing=1.3)
    
    add_p(doc, "Kết thúc", bold=True, underline=True, space_after=8, size=11)
    add_p(doc, "Chủ sự", bold=True, italic=True, space_after=3, size=11)
    add_p(doc, "Xin Thiên Chúa toàn năng, cho ta qua một đêm yên ổn\nVà giờ sau hết được chết lành.", space_after=14, size=11, line_spacing=1.3)
    
    add_p(doc, "Cộng đoàn", bold=True, space_after=3, size=11)
    add_p(doc, "A- men", space_after=8, size=11)
    
    doc.add_page_break() # -> TRANG 41
    
    # =========================================================================
    # TRANG 41: PHỤ LỤC I: KINH NGUYỆN PHANXICÔ - 1. KINH HÒA BÌNH
    # =========================================================================
    add_heading_1(doc, "PHỤ LỤC I: KINH NGUYỆN PHANXICÔ", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=3)
    add_p(doc, "1. KINH HÒA BÌNH", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=1)
    add_p(doc, "(Kinh nguyện của Thánh Phanxicô Assisi)", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, space_after=10)
    
    # Khổ 1
    add_p(doc, "Lạy Chúa từ nhân,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "xin cho con biết mến yêu và phụng sự Chúa trong mọi người.", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Lạy Chúa, xin hãy dùng con như khí cụ bình an của Chúa:", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, size=10)
    
    # Khổ 2 (Đối lập)
    add_p(doc, "Để con đem yêu thương vào nơi oán thù,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Đem thứ tha vào nơi lăng nhục,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Đem an hòa vào nơi tranh chấp,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Đem chân lý vào nơi sai lầm,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Đem đức tin vào nơi nghi nan,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Đem hy vọng vào nơi thất vọng,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Đem ánh sáng vào nơi tối tăm,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Đem niềm vui vào nơi sầu muộn.", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, size=10)
    
    # Khổ 3
    add_p(doc, "Lạy Chúa, xin dạy con:", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Tìm an ủi người hơn được người an ủi,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Tìm hiểu biết người hơn được người hiểu biết,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Tìm yêu mến người hơn được người yêu mến.", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, size=10)
    
    # Khổ 4 (Kết)
    add_p(doc, "Vì chính khi hiến thân là khi được nhận lãnh,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Chính khi thứ tha là khi được tha thứ,", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=10)
    add_p(doc, "Chính khi chết đi là khi vui sống muôn đời.", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3, size=10)
    add_p(doc, "Amen.", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, size=10)
    
    doc.add_page_break() # -> TRANG 42
    
    # =========================================================================
    # TRANG 42: PHỤ LỤC I (tiếp) - 2. KINH CẦU THÁNH PHANXICÔ CHO HÒA BÌNH
    # =========================================================================
    add_heading_1(doc, "2. KINH CẦU THÁNH PHANXICÔ CHO HÒA BÌNH", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=2)
    add_p(doc, "(Kỷ niệm 800 năm Thánh Phanxicô gặp Chị Chết 1226 – 2026)", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, space_after=8)
    
    poem_lines = [
        "Lạy Thánh Phanxicô,/",
        "người anh em của nhân loại,/",
        "cách đây tám trăm năm/",
        "Ngài đã đến gặp Chị Chết với tấm lòng bình an./",
        "Xin chuyển cầu cho chúng con trước nhan thánh Chúa.",
        "Ngài đã nhận ra bình an đích thực/",
        "nơi Đấng chịu đóng đinh trên thánh giá San Damiano,/",
        "Xin dạy chúng con biết tìm nơi Đức Kitô/",
        "suối nguồn hòa giải vốn phá tan mọi bức tường ngăn cách.",
        "Ngài không mang khí giới,/",
        "nhưng dám vượt qua những làn ranh/",
        "của chiến tranh và hiểu lầm./",
        "Xin ban cho chúng con lòng can đảm/",
        "để xây dựng những nhịp cầu/",
        "ở chính nơi thế gian dựng lên các hàng rào chia cắt./",
        "Trong thời đại bị tổn thương sâu sắc/",
        "bởi xung đột và chia rẽ này,/",
        "xin chuyển cầu cùng Chúa cho chúng con/",
        "trở thành những người kiến tạo hòa bình,/",
        "những chứng nhân không cần vũ khí/",
        "biết làm cho thế giới buông bỏ bạo lực/",
        "và làm chứng cho bình an đích thực đến từ Đức Kitô.",
        "Amen."
    ]
    for line in poem_lines:
        add_p(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1.5, size=9.5)
    
    add_p(doc, "ĐGH Leo XIV", bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=2, size=9.5)
    
    doc.add_page_break() # -> TRANG 43
    
    # =========================================================================
    # TRANG 43: PHỤ LỤC II - CÁC CA VÃN KÍNH ĐỨC MẸ (Bài 1, 2) - 2 CỘT
    # =========================================================================
    add_heading_1(doc, "PHỤ LỤC II: CÁC CA VÃN KÍNH ĐỨC MẸ", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=2)
    add_p(doc, "(Hát hoặc đọc vào cuối Giờ Kinh Tối theo mùa phụng vụ)", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=8.5, space_after=6)
    
    table43 = doc.add_table(rows=1, cols=2)
    table43.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr43 = table43._tbl.tblPr
    tblBorders43 = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    tblPr43.append(tblBorders43)
    
    col_w = Mm(58)
    table43.columns[0].width = col_w
    table43.columns[1].width = col_w
    
    cell43_l = table43.cell(0, 0)
    cell43_r = table43.cell(0, 1)
    cell43_l.width = col_w
    cell43_r.width = col_w
    
    # Column 1: Bài 1 (Salve Regina)
    p_l1 = cell43_l.paragraphs[0]
    p_l1.paragraph_format.space_before = Pt(0)
    p_l1.paragraph_format.space_after = Pt(1)
    p_l1.paragraph_format.line_spacing = 1.1
    r_l1 = p_l1.add_run("1. Kính chào Đức Nữ Vương\n(Salve Regina)")
    r_l1.bold = True
    r_l1.font.name = 'Times New Roman'
    r_l1.font.size = Pt(9.5)
    
    p_note1 = cell43_l.add_paragraph()
    p_note1.paragraph_format.space_before = Pt(0)
    p_note1.paragraph_format.space_after = Pt(2)
    r_note1 = p_note1.add_run("(Hát trong Mùa Thường Niên)")
    r_note1.italic = True
    r_note1.font.name = 'Times New Roman'
    r_note1.font.size = Pt(8)
    
    lines_43_1 = [
        "Kính Chào Đức Nữ Vương,",
        "Bà là Mẹ xót thương,",
        "Ngọt ngào cho cuộc sống",
        "Kính chào Lẽ Cậy Trông!*",
        "Này con cháu E-và",
        "Thân phận người lưu lạc,",
        "Chúng con ngửa trông Bà,",
        "Kêu Bà mà khóc lóc,",
        "Than thở với rên la",
        "Trong lũng đầy nước mắt.*",
        "Bà là Nữ Trạng Sư",
        "Nguyện đưa mắt nhân từ,",
        "Phía đoàn con đoái lại;*",
        "Và sau đời khổ ải,",
        "Xin Bà khứng tỏ ra",
        "Cho đoàn con được thấy",
        "Quả phúc bởi lòng Bà:",
        "Đức Giêsu khả ái.*",
        "Ôi lượng cả khoan hồng,",
        "Ôi tấm lòng xót thương,",
        "Ôi dịu hiền nhân hậu,",
        "Trinh nữ Maria."
    ]
    for l in lines_43_1:
        p_c = cell43_l.add_paragraph()
        p_c.paragraph_format.space_before = Pt(0)
        p_c.paragraph_format.space_after = Pt(1.5)
        p_c.paragraph_format.line_spacing = 1.1
        r = p_c.add_run(l)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9)
        
    # Column 2: Bài 2 (Alma Redemptoris Mater)
    p_r1 = cell43_r.paragraphs[0]
    p_r1.paragraph_format.space_before = Pt(0)
    p_r1.paragraph_format.space_after = Pt(1)
    p_r1.paragraph_format.line_spacing = 1.1
    r_r1 = p_r1.add_run("2. Lạy Mẫu Nghi cao cả\n(Alma Redemptoris)")
    r_r1.bold = True
    r_r1.font.name = 'Times New Roman'
    r_r1.font.size = Pt(9.5)
    
    p_note2 = cell43_r.add_paragraph()
    p_note2.paragraph_format.space_before = Pt(0)
    p_note2.paragraph_format.space_after = Pt(2)
    r_note2 = p_note2.add_run("(Hát Mùa Vọng & Giáng Sinh)")
    r_note2.italic = True
    r_note2.font.name = 'Times New Roman'
    r_note2.font.size = Pt(8)
    
    lines_43_2 = [
        "Lạy Mẫu Nghi cao cả,",
        "Sinh dưỡng Chúa cứu đời",
        "Là Cửa Trời rộng mở",
        "Ngôi Sao Biển rạng ngời*",
        "Xin Mẹ thương cứu trợ",
        "Kẻ lỡ bước sa chân",
        "Đang tìm tay nâng đỡ",
        "Mà cải quá tự tân.*",
        "Mẹ sinh Chúa Thiên Đình",
        "Đấng tạo thành nên Mẹ,",
        "Trước sau vẫn khiết trinh,",
        "Ôi lạ lùng khôn ví!*",
        "Gáp-ri-en mừng hát,",
        "Xin Mẹ nhận lời chào,",
        "Và dủ tình thương xót",
        "Đoàn tội lỗi quỳ tâu."
    ]
    for l in lines_43_2:
        p_c = cell43_r.add_paragraph()
        p_c.paragraph_format.space_before = Pt(0)
        p_c.paragraph_format.space_after = Pt(1.5)
        p_c.paragraph_format.line_spacing = 1.1
        r = p_c.add_run(l)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9)
        
    doc.add_page_break() # -> TRANG 44
    
    # =========================================================================
    # TRANG 44: PHỤ LỤC II (tiếp theo) - CA VÃN ĐỨC MẸ (Bài 3, 4, 5) - 2 CỘT
    # =========================================================================
    table44 = doc.add_table(rows=1, cols=2)
    table44.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr44 = table44._tbl.tblPr
    tblBorders44 = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    tblPr44.append(tblBorders44)
    
    table44.columns[0].width = col_w
    table44.columns[1].width = col_w
    
    cell44_l = table44.cell(0, 0)
    cell44_r = table44.cell(0, 1)
    cell44_l.width = col_w
    cell44_r.width = col_w
    
    # Left column: Bài 3 và Bài 4
    p_3 = cell44_l.paragraphs[0]
    p_3.paragraph_format.space_before = Pt(0)
    p_3.paragraph_format.space_after = Pt(1)
    p_3.paragraph_format.line_spacing = 1.1
    r_3 = p_3.add_run("3. Kính lạy Bà, Vị Nữ Hoàng Thiên Quốc\n(Ave Regina caelorum)")
    r_3.bold = True
    r_3.font.name = 'Times New Roman'
    r_3.font.size = Pt(9.5)
    
    p_note3 = cell44_l.add_paragraph()
    p_note3.paragraph_format.space_before = Pt(0)
    p_note3.paragraph_format.space_after = Pt(2)
    r_note3 = p_note3.add_run("(Hát trong Mùa Chay)")
    r_note3.italic = True
    r_note3.font.name = 'Times New Roman'
    r_note3.font.size = Pt(8)
    
    lines_44_3 = [
        "Kính lạy Bà, vị Nữ Hoàng Thiên Quốc,",
        "Lạy Nữ Vương trên chín phẩm thiên thần,",
        "Là Cội Thiêng là Cửa Trời vinh phúc",
        "Đem Vầng Hồng rực rỡ xuống trần gian*",
        "Mừng vui lên, mừng vui lên Trinh Nữ,",
        "Bà hiển vinh, Bà diễm lệ khôn tày.",
        "Bên toà Chúa Kitô, Ngôi Thánh Tử,",
        "Cúi lạy Bà, xin nguyện giúp cầu thay."
    ]
    for l in lines_44_3:
        p_c = cell44_l.add_paragraph()
        p_c.paragraph_format.space_before = Pt(0)
        p_c.paragraph_format.space_after = Pt(1.5)
        p_c.paragraph_format.line_spacing = 1.1
        r = p_c.add_run(l)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9)
        
    p_4 = cell44_l.add_paragraph()
    p_4.paragraph_format.space_before = Pt(6)
    p_4.paragraph_format.space_after = Pt(1)
    p_4.paragraph_format.line_spacing = 1.1
    r_4 = p_4.add_run("4. Lạy Đức Mẹ Chúa Trời\n(Sub tuum praesidium)")
    r_4.bold = True
    r_4.font.name = 'Times New Roman'
    r_4.font.size = Pt(9.5)
    
    p_note4 = cell44_l.add_paragraph()
    p_note4.paragraph_format.space_before = Pt(0)
    p_note4.paragraph_format.space_after = Pt(2)
    r_note4 = p_note4.add_run("(Kinh cổ kính kính Đức Mẹ)")
    r_note4.italic = True
    r_note4.font.name = 'Times New Roman'
    r_note4.font.size = Pt(8)
    
    lines_44_4 = [
        "Lạy Đức Mẹ Chúa Trời,",
        "Ngài xiết bao thánh thiện,",
        "Này chúng con chạy đến",
        "Tìm nương ẩn nơi Ngài.*",
        "Lúc sa vòng gian khổ,",
        "Khi gặp cảnh phong trần,",
        "Lời con cái nài van,",
        "Xin Mẹ đừng chê bỏ.*",
        "Nhưng xin hằng giải thoát,",
        "Khỏi ngàn nỗi hiểm nguy,",
        "Ôi vinh diệu ái bi",
        "Trinh Nữ đầy ơn phước!"
    ]
    for l in lines_44_4:
        p_c = cell44_l.add_paragraph()
        p_c.paragraph_format.space_before = Pt(0)
        p_c.paragraph_format.space_after = Pt(1.5)
        p_c.paragraph_format.line_spacing = 1.1
        r = p_c.add_run(l)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9)
        
    # Right column: Bài 5
    p_5 = cell44_r.paragraphs[0]
    p_5.paragraph_format.space_before = Pt(0)
    p_5.paragraph_format.space_after = Pt(1)
    p_5.paragraph_format.line_spacing = 1.1
    r_5 = p_5.add_run("5. Mừng vui lên, lạy Nữ Hoàng Thiên Quốc\n(Regina caeli)")
    r_5.bold = True
    r_5.font.name = 'Times New Roman'
    r_5.font.size = Pt(9.5)
    
    p_note5 = cell44_r.add_paragraph()
    p_note5.paragraph_format.space_before = Pt(0)
    p_note5.paragraph_format.space_after = Pt(2)
    r_note5 = p_note5.add_run("(Hát trong Mùa Phục Sinh)")
    r_note5.italic = True
    r_note5.font.name = 'Times New Roman'
    r_note5.font.size = Pt(8)
    
    lines_44_5 = [
        "Mừng vui lên, lạy Nữ Hoàng Thiên Quốc",
        "Ha-lê-lui-a!",
        "Vì Thánh Tử Bà được phúc cưu mang",
        "Ha-lê-lui-a!",
        "Đã phục sinh như lời Người phán trước",
        "Ha-lê-lui-a!",
        "Cầu Chúa cho đoàn con, lạy Nữ Hoàng",
        "Ha-lê-lui-a!"
    ]
    for l in lines_44_5:
        p_c = cell44_r.add_paragraph()
        p_c.paragraph_format.space_before = Pt(0)
        p_c.paragraph_format.space_after = Pt(2)
        p_c.paragraph_format.line_spacing = 1.15
        r = p_c.add_run(l)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9)
        if "Ha-lê-lui-a" in l:
            r.italic = True
            
    output_path = "Sach_Kinh_Phung_Vu_Luu_Xa_Phanxico.docx"
    doc.save(output_path)
    print(f"Saved document to {output_path}")

if __name__ == '__main__':
    build_all()
